#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Étape de fusion des PDF après 'partenaires connus', puis export Excel basé sur le modèle.

Déclenchement attendu :
depuis partco_stage._finish_window(app), bouton OK => merge_stage.begin(app, params)

Excel :
- En-tête d'impression : Lieu, Domaine, Rayon, Période
- Ligne 1 : en-têtes du modèle grilledebase.xlsx
- Données à partir de la ligne 2
- Ligne 1 répétée en haut de chaque page imprimée
- A4 paysage
- Toutes les colonnes A:S sur 1 page en largeur
- Hauteur automatique sur plusieurs pages
- PDF tableau A:S séparé
- PDF d'impression séparé : tableau tourné en portrait puis offres
- Dossiers et fichiers finaux avec noms lisibles (ville + domaine)
"""

import csv
import os
import re
import shutil
import unicodedata
import tkinter as tk
from tkinter import ttk, messagebox

from interface import Params, BG
import finalisation as fin

# PyPDF (pypdf ou PyPDF2)
HAS_PYPDF = True
try:
    from pypdf import PdfReader, PdfWriter
except Exception:
    try:
        from PyPDF2 import PdfReader, PdfWriter  # type: ignore
    except Exception:
        HAS_PYPDF = False

# Excel (openpyxl)
HAS_OXL = True
try:
    import openpyxl
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
except Exception:
    HAS_OXL = False

# PDF du tableau (ReportLab)
HAS_REPORTLAB = True
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, KeepTogether
    )
except Exception:
    HAS_REPORTLAB = False


# ---------------------------------------------------------------------------
# Utilitaires généraux
# ---------------------------------------------------------------------------

def _read_list_any(path: str) -> list[str]:
    """Lit JSON (liste) ou une valeur par ligne."""
    try:
        if os.path.exists(path):
            import json
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list):
                    return [str(x) for x in data]
    except Exception:
        pass

    vals = []
    try:
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                s = line.strip()
                if s:
                    vals.append(s)
    except Exception:
        pass
    return vals


def _compute_csv_path(params: Params) -> str:
    return os.path.join(
        params.output_dir or os.getcwd(),
        f"etud-{params.lieux}-{params.emission}jours-{params.domaine}-rayon{params.rayon}km.csv"
    )


def _compute_pdf_dir(params: Params) -> str:
    """
    Retourne le dossier PDF avec un nom lisible.

    Exemple :
        pdf-56121-1jours-H-rayon10km
    devient :
        pdf-lorient-Industrie-1jours-rayon10km

    Si l'ancien dossier existe déjà, il est automatiquement migré.
    """
    base = params.output_dir or os.getcwd()

    readable_dir = os.path.join(base, f"pdf-{_readable_study_stem(params)}")
    legacy_dir = os.path.join(
        base,
        f"pdf-{params.lieux}-{params.emission}jours-{params.domaine}-rayon{params.rayon}km"
    )

    # Étude déjà au nouveau format
    if os.path.isdir(readable_dir):
        # Si une étape précédente a recréé l'ancien dossier, rapatrier son contenu.
        if os.path.isdir(legacy_dir) and os.path.abspath(legacy_dir) != os.path.abspath(readable_dir):
            _merge_directory_contents(legacy_dir, readable_dir)
        _migrate_legacy_aggregate_names(params, readable_dir)
        return readable_dir

    # Migration automatique d'une étude existante
    if os.path.isdir(legacy_dir) and os.path.abspath(legacy_dir) != os.path.abspath(readable_dir):
        try:
            os.rename(legacy_dir, readable_dir)
        except Exception:
            os.makedirs(readable_dir, exist_ok=True)
            _merge_directory_contents(legacy_dir, readable_dir)
    else:
        os.makedirs(readable_dir, exist_ok=True)

    _migrate_legacy_aggregate_names(params, readable_dir)
    return readable_dir


def _sanitize_filename(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*]', "_", name)


def _add_bookmark(writer, title: str, page_index: int):
    """Compatibilité pypdf / PyPDF2."""
    try:
        add = getattr(writer, "add_outline_item", None)
        if add:
            add(title, page_index)
            return
        add2 = getattr(writer, "addBookmark", None)
        if add2:
            add2(title, page_index)
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Résolution des noms lisibles pour l'entête Excel
# ---------------------------------------------------------------------------

def _candidate_paths(params: Params, attr_name: str, filename: str) -> list[str]:
    """
    Construit une liste de chemins possibles pour communes.csv,
    departement.csv et domainepro.csv.
    """
    candidates = []

    # Chemin éventuellement mémorisé dans Params
    p = getattr(params, attr_name, "") or ""
    if p:
        candidates.append(p)

    # Dossier du programme
    try:
        candidates.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), filename))
    except Exception:
        pass

    # Répertoire courant
    candidates.append(os.path.join(os.getcwd(), filename))

    # Dossier de sortie de l'étude
    if getattr(params, "output_dir", ""):
        candidates.append(os.path.join(params.output_dir, filename))

    # Éliminer doublons en gardant l'ordre
    unique = []
    seen = set()
    for p in candidates:
        p = os.path.abspath(p)
        if p not in seen:
            seen.add(p)
            unique.append(p)
    return unique


def _read_csv_rows(path: str) -> list[dict]:
    """Lit un CSV avec plusieurs encodages possibles."""
    last_error = None
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                return list(csv.DictReader(f))
        except Exception as e:
            last_error = e
    if last_error:
        raise last_error
    return []


def _resolve_lieu_label(params: Params) -> str:
    """
    Transforme params.lieux en nom lisible.
    - Commune : code INSEE -> nom_standard
    - Département : ex. 56D -> LIBELLE
    """
    code = str(getattr(params, "lieux", "") or "").strip()
    if not code:
        return "Non précisé"

    is_dep = code.upper().endswith("D")
    raw_code = code[:-1] if is_dep else code

    if is_dep:
        for path in _candidate_paths(params, "departements_csv", "departement.csv"):
            if not os.path.exists(path):
                continue
            try:
                rows = _read_csv_rows(path)
                for row in rows:
                    dep = str(row.get("DEP", "") or "").strip()
                    if dep == raw_code:
                        # colonne LIBELLE ou LIBELLE*
                        label = ""
                        for k, v in row.items():
                            if str(k).strip().upper().startswith("LIBELLE"):
                                label = str(v or "").strip()
                                if label:
                                    break
                        if label:
                            return f"{label} ({raw_code})"
            except Exception:
                pass
        return f"Département {raw_code}"

    for path in _candidate_paths(params, "communes_csv", "communes.csv"):
        if not os.path.exists(path):
            continue
        try:
            rows = _read_csv_rows(path)
            for row in rows:
                c = str(row.get("code_insee", "") or "").strip()
                if c == raw_code:
                    label = str(row.get("nom_standard", "") or "").strip()
                    if label:
                        return f"{label} ({raw_code})"
        except Exception:
            pass

    return raw_code


def _resolve_domaine_label(params: Params) -> str:
    """Transforme le code domaine en libellé lisible via domainepro.csv."""
    code = str(getattr(params, "domaine", "") or "").strip()
    if not code or code.lower() == "tout":
        return "Tous les domaines"

    for path in _candidate_paths(params, "domaines_csv", "domainepro.csv"):
        if not os.path.exists(path):
            continue
        try:
            rows = _read_csv_rows(path)
            for row in rows:
                c = str(row.get("codedomainepro", "") or "").strip()
                if c == code:
                    label = str(row.get("domainepro", "") or "").strip()
                    if label:
                        return f"{label} ({code})"
        except Exception:
            pass

    return code


def _resolve_emission_label(params: Params) -> str:
    """Transforme la valeur emission en période lisible."""
    emission = str(getattr(params, "emission", "") or "").strip().lower()
    mapping = {
        "1": "1 jour",
        "3": "3 jours",
        "7": "1 semaine",
        "14": "2 semaines",
        "31": "1 mois",
        "tout": "Toutes les offres",
        "": "Toutes les offres",
    }
    return mapping.get(emission, f"{emission} jours")


def _study_header_values(params: Params) -> tuple[str, str, str, str]:
    lieu = _resolve_lieu_label(params)
    domaine = _resolve_domaine_label(params)
    rayon = str(getattr(params, "rayon", "") or "").strip()
    if rayon:
        rayon = f"{rayon} km"
    else:
        rayon = "Non précisé"
    periode = _resolve_emission_label(params)
    return lieu, domaine, rayon, periode


# ---------------------------------------------------------------------------
# Noms lisibles des dossiers et fichiers
# ---------------------------------------------------------------------------

def _remove_code_suffix(label: str) -> str:
    """Retire un suffixe technique du type ' (56121)' ou ' (H)'."""
    return re.sub(r"\s*\([^()]*\)\s*$", "", str(label or "")).strip()


def _safe_name_component(value: str, *, lower: bool = False) -> str:
    """
    Nettoie un élément de nom de fichier :
    - enlève les accents ;
    - remplace espaces et séparateurs gênants par '-';
    - conserve les majuscules du libellé sauf si lower=True.
    """
    s = str(value or "").strip()
    if not s:
        s = "inconnu"

    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))

    # Caractères interdits / gênants sous Windows
    s = re.sub(r'[<>:"/\\|?*]+', "-", s)
    s = re.sub(r"[\s_]+", "-", s)
    s = re.sub(r"-{2,}", "-", s)
    s = s.strip(" .-")

    if lower:
        s = s.lower()

    return s or "inconnu"


def _readable_study_stem(params: Params) -> str:
    """
    Partie commune des noms lisibles.

    Exemple :
        lorient-Industrie-1jours-rayon10km
    """
    lieu = _remove_code_suffix(_resolve_lieu_label(params))
    domaine = _remove_code_suffix(_resolve_domaine_label(params))

    lieu_part = _safe_name_component(lieu, lower=True)
    domaine_part = _safe_name_component(domaine)

    emission = str(getattr(params, "emission", "") or "").strip()
    if not emission or emission.lower() == "tout":
        periode_part = "toutes-offres"
    else:
        periode_part = f"{_safe_name_component(emission)}jours"

    rayon = str(getattr(params, "rayon", "") or "0").strip()
    rayon_part = f"rayon{_safe_name_component(rayon)}km"

    return f"{lieu_part}-{domaine_part}-{periode_part}-{rayon_part}"


def _legacy_study_stem(params: Params) -> str:
    return (
        f"{params.lieux}-{params.emission}jours-"
        f"{params.domaine}-rayon{params.rayon}km"
    )


def _merge_directory_contents(source_dir: str, destination_dir: str) -> None:
    """
    Rapatrie le contenu de l'ancien dossier vers le nouveau.
    Les fichiers déjà présents dans le nouveau dossier sont conservés.
    """
    try:
        os.makedirs(destination_dir, exist_ok=True)
        for name in os.listdir(source_dir):
            src = os.path.join(source_dir, name)
            dst = os.path.join(destination_dir, name)

            if os.path.exists(dst):
                continue

            try:
                shutil.move(src, dst)
            except Exception:
                pass

        try:
            if not os.listdir(source_dir):
                os.rmdir(source_dir)
        except Exception:
            pass
    except Exception:
        pass


def _migrate_legacy_aggregate_names(params: Params, pdf_dir: str) -> None:
    """Renomme les anciens fichiers globaux s'ils existent déjà."""
    legacy = _legacy_study_stem(params)
    readable = _readable_study_stem(params)

    pairs = [
        (f"fusion-{legacy}.pdf", f"fusion-{readable}.pdf"),
        (f"fusion-{legacy}.xlsx", f"fusion-{readable}.xlsx"),
        (f"tableau-{legacy}.pdf", f"tableau-{readable}.pdf"),
    ]

    for old_name, new_name in pairs:
        old_path = os.path.join(pdf_dir, old_name)
        new_path = os.path.join(pdf_dir, new_name)

        if not os.path.exists(old_path) or os.path.exists(new_path):
            continue

        try:
            os.rename(old_path, new_path)
        except Exception:
            pass


def _ensure_readable_csv_copy(params: Params, csv_path: str) -> str:
    """
    Conserve le CSV technique utilisé par les étapes précédentes et crée,
    en plus, une copie avec un nom lisible.
    """
    if not csv_path or not os.path.exists(csv_path):
        return ""

    readable_path = os.path.join(
        params.output_dir or os.getcwd(),
        f"etud-{_readable_study_stem(params)}.csv"
    )

    if os.path.abspath(readable_path) == os.path.abspath(csv_path):
        return readable_path

    try:
        shutil.copy2(csv_path, readable_path)
        return readable_path
    except Exception:
        return ""


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------

def _export_excel_from_df(params: Params) -> str:
    """
    Construit l'Excel à partir du CSV et du modèle grilledebase.xlsx.
    """
    if not HAS_OXL:
        messagebox.showerror(
            "Export Excel",
            "openpyxl n'est pas disponible. Installe 'openpyxl'."
        )
        return ""

    csv_path = _compute_csv_path(params)
    pdf_dir = _compute_pdf_dir(params)
    out_xlsx = os.path.join(
        pdf_dir,
        f"fusion-{_readable_study_stem(params)}.xlsx"
    )

    # Charger le CSV
    df = fin._load_df(csv_path)
    if df.empty or df.shape[1] < 5:
        messagebox.showerror(
            "Export Excel",
            "Le CSV est vide ou invalide (moins de 5 colonnes)."
        )
        return ""

    # Ignorer la première ligne du DataFrame comme dans la version existante
    if len(df) >= 1:
        df_data = df.iloc[1:].reset_index(drop=True)
    else:
        df_data = df.copy()

    # Chercher le modèle
    template_candidates = [
        os.path.join(os.getcwd(), "grilledebase.xlsx"),
        os.path.join(
            os.path.dirname(__file__) if "__file__" in globals() else os.getcwd(),
            "grilledebase.xlsx"
        ),
        os.path.join(os.path.dirname(csv_path), "grilledebase.xlsx"),
    ]

    template_path = ""
    for p in template_candidates:
        if os.path.exists(p):
            template_path = p
            break

    if not template_path:
        messagebox.showerror(
            "Export Excel",
            "Modèle 'grilledebase.xlsx' introuvable. "
            "Place-le à la racine du projet."
        )
        return ""

    # Charger le modèle
    try:
        wb = load_workbook(template_path)
    except Exception as e:
        messagebox.showerror(
            "Export Excel",
            f"Impossible d'ouvrir le modèle Excel :\n{e}"
        )
        return ""

    ws = wb.active

    # -----------------------------------------------------------------------
    # Informations de l'étude dans le véritable en-tête d'impression Excel.
    # Aucune ligne n'est ajoutée dans la feuille : la ligne 1 du modèle reste
    # la ligne des titres du tableau.
    # -----------------------------------------------------------------------
    lieu, domaine, rayon, periode = _study_header_values(params)

    try:
        ws.oddHeader.left.text = f"Lieu : {lieu}\nRayon : {rayon}"
        ws.oddHeader.center.text = f"Domaine : {domaine}"
        ws.oddHeader.right.text = f"Période : {periode}"

        ws.evenHeader.left.text = f"Lieu : {lieu}\nRayon : {rayon}"
        ws.evenHeader.center.text = f"Domaine : {domaine}"
        ws.evenHeader.right.text = f"Période : {periode}"

        for section in (
            ws.oddHeader.left, ws.oddHeader.center, ws.oddHeader.right,
            ws.evenHeader.left, ws.evenHeader.center, ws.evenHeader.right,
        ):
            section.size = 9
            section.font = "Arial,Bold"
    except Exception:
        pass

    # Les données commencent comme avant en ligne 2.
    # La ligne 1 reste la ligne des en-têtes du tableau.
    start_row = 2
    n = len(df_data)

    for i in range(n):
        row_excel = start_row + i
        row = df_data.iloc[i]

        # A : index
        try:
            a_val = int(str(row.iloc[0]).strip())
        except Exception:
            try:
                a_val = float(str(row.iloc[0]).strip())
            except Exception:
                a_val = row.iloc[0]

        ws.cell(row=row_excel, column=1, value=a_val)

        # B : référence
        ws.cell(row=row_excel, column=2, value=str(row.iloc[1]))

        # C : partenaire
        ws.cell(row=row_excel, column=3, value=str(row.iloc[2]))

        # D : type de contrat (df col 5)
        ws.cell(row=row_excel, column=4, value=row.iloc[4])

        # E:R vides
        for col in range(5, 19):
            ws.cell(row=row_excel, column=col, value=None)

        # S
        formula = f"=COUNTIF(F{row_excel}:R{row_excel},1)>0"
        ws.cell(row=row_excel, column=19, value=formula)

    # -----------------------------------------------------------------------
    # Mise en forme lignes de données
    # -----------------------------------------------------------------------
    if n > 0:
        data_first = start_row
        data_last = start_row + n - 1

        zebra_fill = PatternFill(
            fill_type="solid",
            start_color="FFD9D9D9",
            end_color="FFD9D9D9"
        )
        thin = Side(style="thin")
        thin_border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin
        )

        for r in range(data_first, data_last + 1):
            is_grey = ((r - data_first) % 2 == 0)
            for c in range(1, 20):
                cell = ws.cell(row=r, column=c)
                if is_grey:
                    cell.fill = zebra_fill
                cell.border = thin_border

    # -----------------------------------------------------------------------
    # Ligne des totaux
    # -----------------------------------------------------------------------
    if n > 0:
        total_row = start_row + n

        ws.cell(row=total_row, column=1, value="total")
        ws.cell(row=total_row, column=2, value=None)
        ws.cell(row=total_row, column=3, value=None)
        ws.cell(row=total_row, column=4, value=None)

        for col in range(5, 19):
            col_letter = get_column_letter(col)
            ws.cell(
                row=total_row,
                column=col,
                value=f"=SUM({col_letter}{start_row}:{col_letter}{total_row-1})"
            )

        ws.cell(
            row=total_row,
            column=19,
            value=f"=COUNTIF(S{start_row}:S{total_row-1},TRUE)"
        )

        thin = Side(style="thin")
        medium = Side(style="medium")

        for c in range(1, 20):
            cell = ws.cell(row=total_row, column=c)
            cell.border = Border(
                left=thin,
                right=thin,
                top=medium,
                bottom=thin
            )

        last_row = total_row
    else:
        last_row = start_row

    # -----------------------------------------------------------------------
    # Mise en page impression
    # -----------------------------------------------------------------------
    try:
        # Répéter uniquement la ligne des titres du tableau
        ws.print_title_rows = "1:1"

        # Zone d'impression
        ws.print_area = f"A1:S{last_row}"

        # Activer réellement "Ajuster à"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.pageSetUpPr.autoPageBreaks = False

        # A4 paysage, 1 page en largeur, hauteur libre
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.scale = None

        # Marges réduites
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.65
        ws.page_margins.bottom = 0.35
        ws.page_margins.header = 0.20
        ws.page_margins.footer = 0.15

        ws.print_options.horizontalCentered = True

    except Exception:
        pass

    # Recalcul à l'ouverture
    try:
        wb.calc_properties.fullCalcOnLoad = True
    except Exception:
        pass

    # Écraser le fichier s'il existe
    try:
        if os.path.exists(out_xlsx):
            os.remove(out_xlsx)
    except Exception:
        pass

    try:
        wb.save(out_xlsx)
    except Exception as e:
        messagebox.showerror(
            "Export Excel",
            f"Erreur lors de l'enregistrement du fichier Excel :\n{e}"
        )
        return ""

    return out_xlsx



# ---------------------------------------------------------------------------
# PDF du tableau + PDF d'impression (tableau tourné en portrait puis offres)
# ---------------------------------------------------------------------------

def _compute_xlsx_path(params: Params) -> str:
    pdf_dir = _compute_pdf_dir(params)
    return os.path.join(
        pdf_dir,
        f"fusion-{_readable_study_stem(params)}.xlsx"
    )


def _compute_table_pdf_path(params: Params) -> str:
    pdf_dir = _compute_pdf_dir(params)
    return os.path.join(
        pdf_dir,
        f"tableau-{_readable_study_stem(params)}.pdf"
    )


def _compute_print_pdf_path(params: Params) -> str:
    pdf_dir = _compute_pdf_dir(params)
    return os.path.join(
        pdf_dir,
        f"imprim-{_readable_study_stem(params)}.pdf"
    )


def _pdf_cell_text(value, row_index: int, col_index: int, total_row: bool = False) -> str:
    """Valeur lisible dans le PDF du tableau."""
    if value is None:
        return ""

    s = str(value)

    # Ne pas imprimer les formules Excel elles-mêmes.
    if s.startswith("="):
        if total_row:
            return "0"
        if col_index == 19:
            return "FAUX"
        return ""

    if isinstance(value, bool):
        return "VRAI" if value else "FAUX"

    return s


def _create_table_pdf_from_excel(params: Params, xlsx_path: str) -> str:
    """
    Génère un PDF A4 paysage du tableau Excel A:S.
    L'en-tête Lieu / Domaine / Rayon / Période est répété sur chaque page.
    """
    if not HAS_REPORTLAB:
        messagebox.showerror(
            "PDF du tableau",
            "ReportLab n'est pas installé.\n\n"
            "Installe-le avec :\n"
            "pip install reportlab"
        )
        return ""

    if not os.path.exists(xlsx_path):
        messagebox.showerror(
            "PDF du tableau",
            f"Fichier Excel introuvable :\n{xlsx_path}"
        )
        return ""

    try:
        wb = load_workbook(xlsx_path, data_only=False)
        ws = wb.active
    except Exception as e:
        messagebox.showerror(
            "PDF du tableau",
            f"Impossible d'ouvrir le fichier Excel :\n{e}"
        )
        return ""

    # Trouver la dernière ligne réellement utilisée dans A:S.
    last_row = 1
    for r in range(ws.max_row, 0, -1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, 20)):
            last_row = r
            break

    # Styles texte.
    styles = getSampleStyleSheet()
    header_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=6.4,
        leading=7.2,
        alignment=TA_CENTER,
        spaceAfter=0,
        spaceBefore=0,
    )
    cell_style = ParagraphStyle(
        "BodyCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=6.2,
        leading=7.0,
        alignment=TA_CENTER,
        spaceAfter=0,
        spaceBefore=0,
    )
    total_style = ParagraphStyle(
        "TotalCell",
        parent=cell_style,
        fontName="Helvetica-Bold",
    )

    table_data = []

    for r in range(1, last_row + 1):
        is_total = str(ws.cell(r, 1).value or "").strip().lower() == "total"
        row_items = []
        for c in range(1, 20):
            value = _pdf_cell_text(
                ws.cell(r, c).value,
                row_index=r,
                col_index=c,
                total_row=is_total,
            )
            # Échapper le minimum pour les Paragraph ReportLab.
            value = (
                value.replace("&", "&amp;")
                     .replace("<", "&lt;")
                     .replace(">", "&gt;")
            )
            style = header_style if r == 1 else (total_style if is_total else cell_style)
            row_items.append(Paragraph(value, style))
        table_data.append(row_items)

    page_w, page_h = landscape(A4)
    left_margin = 18
    right_margin = 18
    top_margin = 62
    bottom_margin = 20
    usable_w = page_w - left_margin - right_margin

    # Reprendre autant que possible les largeurs du fichier Excel, puis
    # les mettre à l'échelle pour tenir exactement sur A:S en A4 paysage.
    weights = []
    for c in range(1, 20):
        letter = get_column_letter(c)
        width = ws.column_dimensions[letter].width
        if width is None:
            width = 8.43
        # Éviter les colonnes quasi invisibles ou exagérément larges.
        weight = min(max(float(width), 4.0), 24.0)
        weights.append(weight)

    # Garantir un peu plus de place aux 4 premières colonnes.
    weights[0] = max(weights[0], 5.0)
    weights[1] = max(weights[1], 9.0)
    weights[2] = max(weights[2], 10.0)
    weights[3] = max(weights[3], 7.0)

    scale = usable_w / sum(weights)
    col_widths = [w * scale for w in weights]

    lieu, domaine, rayon, periode = _study_header_values(params)

    def _draw_page_header(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 8.5)

        y1 = page_h - 21
        y2 = page_h - 33

        canvas.drawString(left_margin, y1, f"Lieu : {lieu}")
        canvas.drawString(left_margin, y2, f"Rayon : {rayon}")

        canvas.drawCentredString(page_w / 2, y1, f"Domaine : {domaine}")
        canvas.drawRightString(page_w - right_margin, y1, f"Période : {periode}")

        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(
            page_w - right_margin,
            9,
            f"Page {doc.page}"
        )
        canvas.restoreState()

    out_pdf = _compute_table_pdf_path(params)

    try:
        if os.path.exists(out_pdf):
            os.remove(out_pdf)
    except Exception:
        pass

    doc = SimpleDocTemplate(
        out_pdf,
        pagesize=landscape(A4),
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin,
        title="Tableau de l'étude d'offres",
        author="CNTPEP",
    )

    table = Table(
        table_data,
        colWidths=col_widths,
        repeatRows=1,
        hAlign="CENTER",
    )

    style_cmds = [
        ("GRID", (0, 0), (-1, -1), 0.35, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 1.2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 1.2),
        ("TOPPADDING", (0, 0), (-1, -1), 2.2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
    ]

    # Zébrage comme dans Excel, à partir de la première ligne de données.
    for r in range(1, len(table_data)):
        if str(ws.cell(r + 1, 1).value or "").strip().lower() == "total":
            style_cmds.append(("LINEABOVE", (0, r), (-1, r), 1.1, colors.black))
            style_cmds.append(("BACKGROUND", (0, r), (-1, r), colors.white))
        elif (r - 1) % 2 == 0:
            style_cmds.append(
                ("BACKGROUND", (0, r), (-1, r), colors.HexColor("#E7E7E7"))
            )

    table.setStyle(TableStyle(style_cmds))

    try:
        doc.build(
            [table],
            onFirstPage=_draw_page_header,
            onLaterPages=_draw_page_header,
        )
    except Exception as e:
        messagebox.showerror(
            "PDF du tableau",
            f"Erreur lors de la création du PDF du tableau :\n{e}"
        )
        return ""

    return out_pdf


def _rotate_page_to_portrait(page):
    """Tourne une page paysage de 90° pour un affichage/impression portrait."""
    try:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
    except Exception:
        width = height = 0

    if width > height:
        try:
            # pypdf récent
            page.rotate(90)
        except Exception:
            try:
                # PyPDF2 ancien
                page.rotateClockwise(90)
            except Exception:
                pass
    return page


def _create_print_pdf(params: Params, table_pdf: str, offers_pdf: str) -> str:
    """
    Crée le PDF d'impression :
    1) pages du tableau, tournées en portrait ;
    2) PDF fusionné des offres, conservé tel quel.
    """
    if not table_pdf or not os.path.exists(table_pdf):
        return ""

    if not offers_pdf or not os.path.exists(offers_pdf):
        messagebox.showerror(
            "PDF d'impression",
            f"Le PDF des offres est introuvable :\n{offers_pdf}"
        )
        return ""

    out_print = _compute_print_pdf_path(params)

    writer = PdfWriter()
    page_offset = 0

    try:
        table_reader = PdfReader(table_pdf)
        _add_bookmark(writer, "Tableau", 0)

        for page in table_reader.pages:
            writer.add_page(_rotate_page_to_portrait(page))

        page_offset = len(table_reader.pages)

        offers_reader = PdfReader(offers_pdf)
        if len(offers_reader.pages) > 0:
            _add_bookmark(writer, "Offres", page_offset)

        for page in offers_reader.pages:
            writer.add_page(page)

        try:
            if os.path.exists(out_print):
                os.remove(out_print)
        except Exception:
            pass

        with open(out_print, "wb") as f:
            writer.write(f)

    except Exception as e:
        messagebox.showerror(
            "PDF d'impression",
            f"Erreur lors de la création du PDF d'impression :\n{e}"
        )
        return ""

    return out_print


def _create_all_outputs(params: Params, offers_pdf: str) -> tuple[str, str, str]:
    """
    Produit l'Excel, le PDF du tableau et le PDF d'impression.
    Retourne (xlsx, tableau_pdf, imprim_pdf).
    """
    out_xlsx = _export_excel_from_df(params)
    if not out_xlsx:
        return "", "", ""

    table_pdf = _create_table_pdf_from_excel(params, out_xlsx)
    if not table_pdf:
        return out_xlsx, "", ""

    print_pdf = _create_print_pdf(params, table_pdf, offers_pdf)
    return out_xlsx, table_pdf, print_pdf


# ---------------------------------------------------------------------------
# Retour à l'interface principale après une étude
# ---------------------------------------------------------------------------

def _return_to_main_window(app, params: Params, print_pdf: str = "") -> None:
    """
    Termine l'étude et reconstruit le véritable écran d'accueil.

    finalisation.py remplace le contenu de la fenêtre principale par l'écran
    « Le fichier existe ». Il faut donc appeler app.reload_main_interface()
    et pas seulement deiconify()/lift().

    - conserve les StringVar et donc les champs saisis ;
    - ne ferme pas Chrome ;
    - reconstruit Domaine / Lieu / Rayon / Dossier ;
    - écrit « Étude terminée » dans le nouveau journal ;
    - remet la fenêtre principale au premier plan.
    """
    study_name = _readable_study_stem(params)

    # Relâcher un éventuel grab posé par une fenêtre intermédiaire.
    try:
        current_grab = app.grab_current()
        if current_grab is not None:
            try:
                current_grab.grab_release()
            except Exception:
                pass
    except Exception:
        pass

    # Détruire proprement une éventuelle mini-fenêtre de finalisation encore référencée.
    try:
        suivant = getattr(app, "_suivant_win", None)
        if suivant is not None and suivant.winfo_exists():
            suivant.destroy()
    except Exception:
        pass
    try:
        app._suivant_win = None
    except Exception:
        pass

    # Reconstruire le VRAI écran initial. Les StringVar ne sont pas recréées,
    # donc les valeurs Domaine / Commune / Rayon / Dossier sont conservées.
    try:
        reload_fn = getattr(app, "reload_main_interface", None)
        if callable(reload_fn):
            reload_fn()
    except Exception:
        pass

    # La fenêtre principale a pu être withdraw() pendant la finalisation.
    try:
        app._main_hidden = False
    except Exception:
        pass
    try:
        app.deiconify()
    except Exception:
        pass

    # Écrire dans le journal APRÈS reload_main_interface(), car l'ancien journal
    # est détruit pendant la reconstruction de l'écran.
    try:
        log = getattr(app, "log", None)
        if log is not None and hasattr(log, "write"):
            log.write(f"Étude terminée : {study_name}")
            if print_pdf:
                log.write(f"Fichier d'impression : {os.path.basename(print_pdf)}")
    except Exception:
        pass

    try:
        app.lift()
        app.focus_force()
        app.update_idletasks()
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Fusion PDF + enchaînement export Excel
# ---------------------------------------------------------------------------

def begin(app, params: Params):
    """
    Démarre l'étape de fusion des PDF puis l'export Excel.
    """
    if not HAS_PYPDF:
        messagebox.showerror(
            "Fusion PDF",
            "Impossible d'importer pypdf/PyPDF2.\n"
            "Installe 'pypdf' ou 'PyPDF2'."
        )
        return

    csv_path = _compute_csv_path(params)

    if not os.path.exists(csv_path):
        messagebox.showerror(
            "CSV introuvable",
            f"Impossible de trouver le CSV :\n{csv_path}\n\n"
            f"Vérifie les paramètres "
            f"(lieux / émission / domaine / rayon / dossier de sortie)."
        )
        return

    # Copie lisible du CSV, sans casser les étapes qui utilisent encore le nom technique.
    _ensure_readable_csv_copy(params, csv_path)

    # À ce moment, l'ancien dossier PDF est automatiquement renommé/migré.
    pdf_dir = _compute_pdf_dir(params)

    out_pdf = os.path.join(
        pdf_dir,
        f"fusion-{_readable_study_stem(params)}.pdf"
    )

    # Si le PDF fusionné existe déjà
    if os.path.exists(out_pdf):
        overwrite = messagebox.askyesno(
            "Fichier fusionné déjà présent",
            "Le fichier de fusion existe déjà.\n\n"
            "Voulez-vous l'écraser et relancer la fusion ?\n\n"
            "Oui = Écraser et refusionner\n"
            "Non = Passer directement à l'export Excel"
        )

        if not overwrite:
            out_xlsx, table_pdf, print_pdf = _create_all_outputs(params, out_pdf)

            created = [p for p in (out_xlsx, table_pdf, print_pdf) if p]
            if created:
                if print_pdf:
                    messagebox.showinfo(
                        "Étude terminée",
                        f"Étude terminée.\n\nFichier d'impression créé :\n{print_pdf}"
                    )
                else:
                    messagebox.showinfo(
                        "Étude terminée",
                        "Étude terminée.\n\nFichiers créés :\n" + "\n".join(created)
                    )

            final_print_pdf = print_pdf if print_pdf else ""
            _return_to_main_window(app, params, final_print_pdf)
            return

        try:
            os.remove(out_pdf)
        except Exception as e:
            messagebox.showerror(
                "Fusion PDF",
                f"Impossible d'écraser le fichier existant :\n"
                f"{out_pdf}\n{e}"
            )
            return

    # Lire listes partenaires
    target_dir = os.path.dirname(csv_path) or os.getcwd()

    partco_set = {
        s.strip().lower()
        for s in _read_list_any(os.path.join(target_dir, "partco.txt"))
        if str(s).strip()
    }

    robot_set = {
        s.strip().lower()
        for s in _read_list_any(os.path.join(target_dir, "robot.txt"))
        if str(s).strip()
    }

    # Charger CSV et construire ordre
    df = fin._load_df(csv_path)

    if df.empty or df.shape[1] < 4:
        messagebox.showerror(
            "Fusion PDF",
            "Le CSV est vide ou invalide (moins de 4 colonnes)."
        )
        return

    merge_items: list[tuple[str, int, str, str, bool]] = []

    for _, row in df.iterrows():
        try:
            idx = int(str(row.iloc[0]).strip())
        except Exception:
            continue

        ref = str(row.iloc[1]).strip()
        partner = str(row.iloc[2]).strip()
        partner_l = partner.lower()
        ref_sane = _sanitize_filename(ref)

        if partner_l == "pe":
            p1 = os.path.join(
                pdf_dir,
                f"{idx:04d}-{ref_sane}.pdf"
            )
            merge_items.append((p1, idx, ref, partner, False))
            continue

        if partner_l not in partco_set:
            continue

        if partner_l in robot_set:
            p1 = os.path.join(
                pdf_dir,
                f"{idx:04d}-{ref_sane}.pdf"
            )
            merge_items.append((p1, idx, ref, partner, False))
        else:
            p1 = os.path.join(
                pdf_dir,
                f"{idx:04d}-{ref_sane}.pdf"
            )
            p2 = os.path.join(
                pdf_dir,
                f"{idx:04d}-par-{ref_sane}.pdf"
            )
            merge_items.append((p1, idx, ref, partner, False))
            merge_items.append((p2, idx, ref, partner, True))

    # Aucun PDF à fusionner : créer au moins Excel + tableau PDF.
    if not merge_items:
        out_xlsx = _export_excel_from_df(params)
        table_pdf = ""

        if out_xlsx:
            table_pdf = _create_table_pdf_from_excel(params, out_xlsx)

        created = [p for p in (out_xlsx, table_pdf) if p]
        if created:
            messagebox.showinfo(
                "Étude terminée",
                "Étude terminée.\n\nFichiers créés :\n" + "\n".join(created)
            )

        _return_to_main_window(app, params, "")
        return

    # Vérifier fichiers manquants
    missing = [
        p for (p, *_rest) in merge_items
        if not os.path.exists(p)
    ]

    if missing:
        msg = (
            "PDF manquant — fusion interrompue :\n"
            + "\n".join(missing[:20])
        )

        if len(missing) > 20:
            msg += f"\n(+ {len(missing)-20} autres…)"

        messagebox.showerror("Fusion PDF", msg)
        return

    # -----------------------------------------------------------------------
    # UI fusion
    # -----------------------------------------------------------------------
    win = tk.Toplevel(app)
    win.title("Fusion des PDF")
    win.configure(bg=BG)
    win.geometry("560x240+0+630")
    win.resizable(False, False)

    tk.Label(
        win,
        text="Fusion des PDF (ordre CSV)",
        bg=BG,
        font=("Arial", 11, "bold")
    ).pack(
        anchor="w",
        padx=12,
        pady=(10, 4)
    )

    frm = ttk.Frame(win, padding=12)
    frm.pack(fill="both", expand=True)

    total = len(merge_items)

    var_prog = tk.StringVar(value=f"0 / {total}")
    var_file = tk.StringVar(value="Fichier : ")
    var_stat = tk.StringVar(value="Statut : prêt.")

    ttk.Label(frm, textvariable=var_prog).pack(anchor="w")
    ttk.Label(frm, textvariable=var_file).pack(anchor="w")
    ttk.Label(
        frm,
        textvariable=var_stat
    ).pack(
        anchor="w",
        pady=(4, 6)
    )

    pbar = ttk.Progressbar(
        frm,
        orient="horizontal",
        mode="determinate",
        length=520
    )
    pbar.pack(anchor="w", pady=(0, 6))
    pbar["maximum"] = total
    pbar["value"] = 0

    writer = PdfWriter()
    page_offset = 0
    state = {"i": 0}

    def _finish_all(merge_ok: bool = True):
        try:
            win.destroy()
        except Exception:
            pass

        if merge_ok and os.path.exists(out_pdf):
            out_xlsx, table_pdf, print_pdf = _create_all_outputs(params, out_pdf)
            created = [p for p in (out_xlsx, table_pdf, print_pdf) if p]
        else:
            out_xlsx = _export_excel_from_df(params)
            table_pdf = _create_table_pdf_from_excel(params, out_xlsx) if out_xlsx else ""
            created = [p for p in (out_xlsx, table_pdf) if p]

        if created:
            if merge_ok and 'print_pdf' in locals() and print_pdf:
                messagebox.showinfo(
                    "Étude terminée",
                    f"Étude terminée.\n\nFichier d'impression créé :\n{print_pdf}"
                )
            else:
                messagebox.showinfo(
                    "Étude terminée",
                    "Étude terminée.\n\nFichiers créés :\n" + "\n".join(created)
                )

        final_print_pdf = print_pdf if (merge_ok and 'print_pdf' in locals() and print_pdf) else ""
        _return_to_main_window(app, params, final_print_pdf)

    def _step():
        nonlocal page_offset

        i = state["i"]

        if i >= total:
            try:
                with open(out_pdf, "wb") as f:
                    writer.write(f)
            except Exception as e:
                messagebox.showerror(
                    "Fusion PDF",
                    f"Erreur lors de l'écriture du PDF fusionné :\n{e}"
                )
                _finish_all(False)
                return

            _finish_all(True)
            return

        path, idx, ref, partner, is_par = merge_items[i]

        var_file.set(f"Fichier : {os.path.basename(path)}")
        var_stat.set("Statut : ajout…")
        win.update_idletasks()

        try:
            reader = PdfReader(path)
            nb = len(reader.pages)

            if is_par:
                title = f"{idx:04d} – {ref} – {partner}"
            else:
                title = f"{idx:04d} – {ref}"

            _add_bookmark(writer, title, page_offset)

            for p in range(nb):
                writer.add_page(reader.pages[p])

            page_offset += nb

        except Exception as e:
            messagebox.showerror(
                "Fusion PDF",
                f"Erreur lecture/ajout :\n{path}\n{e}"
            )
            _finish_all(False)
            return

        pbar["value"] = i + 1
        var_prog.set(f"{i+1} / {total}")
        state["i"] = i + 1

        win.after(200, _step)

    win.after(300, _step)
